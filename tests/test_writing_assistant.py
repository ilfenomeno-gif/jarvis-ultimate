"""
Suite di test unitari per il plugin writing_assistant.
Utilizza pytest e unittest.mock per isolare le dipendenze esterne (es. API DeepSeek, I/O file complessi).
Include test specializzati per la preservazione dei Runs nei file Word (.docx).
"""

import os
import sys
import tempfile
import shutil
from unittest.mock import patch, MagicMock
import pytest

# Assicuriamoci di poter importare il plugin dalla directory root
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from plugins import writing_assistant


@pytest.fixture
def temp_dir():
    """Fornisce una directory temporanea pulita per i file di test."""
    d = tempfile.mkdtemp()
    yield d
    shutil.rmtree(d)


@pytest.fixture
def mock_openai(monkeypatch):
    """Mocka la libreria openai e setta una API key fittizia per evitare chiamate reali."""
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test_key_123")
    
    mock_openai_module = MagicMock()
    mock_client = MagicMock()
    mock_openai_module.OpenAI.return_value = mock_client
    
    mock_response = MagicMock()
    mock_response.choices[0].message.content = "Testo migliorato di test."
    mock_client.chat.completions.create.return_value = mock_response
    
    mock_openai_module.RateLimitError = type('RateLimitError', (Exception,), {})
    mock_openai_module.APIConnectionError = type('APIConnectionError', (Exception,), {})
    mock_openai_module.APIError = type('APIError', (Exception,), {})
    
    monkeypatch.setattr(writing_assistant, "openai", mock_openai_module)
    return mock_client, mock_openai_module


def test_replace_paragraph_text_preserving_runs():
    """Verifica che i run mantengano lo stile e che il testo vi venga redistribuito sopra proporzionalmente."""
    from docx import Document
    doc = Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Grassetto! ")
    r1.bold = True
    r2 = p.add_run("Corsivo!")
    r2.italic = True
    
    # "Grassetto! Corsivo!" = 19 caratteri
    # "Grassetto! " = 11 caratteri (58%)
    # "Corsivo!" = 8 caratteri (42%)
    
    writing_assistant._replace_paragraph_text_preserving_runs(p, "Nuovo Testo Modificato")
    
    # Verifica stili
    assert p.runs[0].bold is True
    assert p.runs[1].italic is True
    
    # Verifica che le parti unificate ricreino il testo atteso
    reconstructed = p.runs[0].text + p.runs[1].text
    assert reconstructed == "Nuovo Testo Modificato"


@patch('plugins.writing_assistant._process_text_with_deepseek', return_value="TestoMigliorato")
def test_process_docx_inplace(mock_deepseek):
    """Verifica che il modulo processi paragrafi, tabelle ed escluda l'alterazione degli stili."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("Paragrafo Standard")
    
    table = doc.add_table(rows=1, cols=1)
    table.cell(0,0).text = "Cella Tabella"
    
    orig, impr = writing_assistant._process_docx_inplace(doc, "", "", None)
    
    # Verifica i testi tornati per i log diff
    assert "Paragrafo Standard\nCella Tabella" == orig
    assert "TestoMigliorato\nTestoMigliorato" == impr
    
    # Verifica che il doc sia stato effettivamente manipolato in-place
    assert doc.paragraphs[0].text == "TestoMigliorato"
    assert doc.tables[0].cell(0,0).text == "TestoMigliorato"


def test_split_into_chunks():
    text = "P1\n\nP2\n\nP3"
    chunks = writing_assistant._split_into_chunks(text, 100)
    assert chunks == ["P1\n\nP2\n\nP3"]
    
    chunks = writing_assistant._split_into_chunks(text, 5)
    assert chunks == ["P1", "P2", "P3"]


def test_extract_frontmatter():
    text_with_fm = "---\ntitle: Test\n---\nContenuto del file."
    fm, clean = writing_assistant._extract_frontmatter(text_with_fm)
    assert fm == "---\ntitle: Test\n---"
    assert clean == "Contenuto del file."


def test_mask_and_unmask_code_blocks():
    text = "Testo normale.\n```python\nprint('hello')\n```\nAltro testo."
    masked, placeholders = writing_assistant._mask_code_blocks(text)
    assert "```python" not in masked
    assert "[[[CODE_BLOCK_0]]]" in masked
    
    unmasked = writing_assistant._unmask_code_blocks(masked, placeholders)
    assert unmasked == text


def test_read_file_txt_md(temp_dir):
    fpath = os.path.join(temp_dir, "test.md")
    with open(fpath, "w", encoding="utf-8") as f:
        f.write("Test content")
        
    content, ext = writing_assistant._read_file(fpath)
    assert content == "Test content"
    assert ext == ".md"


def test_write_file_txt_md(temp_dir):
    fpath = os.path.join(temp_dir, "out.md")
    res_path = writing_assistant._write_file(fpath, "Nuovo", ".md")
    
    with open(fpath, "r", encoding="utf-8") as f:
        assert f.read() == "Nuovo"


def test_generate_diff_log():
    mock_player = MagicMock()
    orig = "Linea A\nLinea B\nLinea C"
    impr = "Linea A\nLinea Modificata\nLinea C\nNuova"
    
    writing_assistant._generate_diff_log(orig, impr, "fake.txt", mock_player)
    assert mock_player.write_log.call_count >= 2


def test_backup_and_cleanup(temp_dir):
    fpath = os.path.join(temp_dir, "test.md")
    with open(fpath, "w", encoding="utf-8") as f:
        f.write("Contenuto")
        
    for i in range(6):
        with open(f"{fpath}.bak_2000010{i}_000000", "w") as f:
            f.write("Old")
            
    with patch("plugins.writing_assistant._process_text_with_deepseek", return_value="Nuovo"):
        writing_assistant.run({"file_path": fpath})
        
    backups = [f for f in os.listdir(temp_dir) if f.startswith("test.md.bak_")]
    assert len(backups) == 5
    

def test_handle_annulla(temp_dir):
    fpath = os.path.join(temp_dir, "test.md")
    with open(fpath, "w", encoding="utf-8") as f:
        f.write("Rovinato")
        
    with open(f"{fpath}.bak_1000", "w", encoding="utf-8") as f:
        f.write("Vecchio")
    with open(f"{fpath}.bak_2000", "w", encoding="utf-8") as f:
        f.write("Originale")
        
    res = writing_assistant._handle_annulla(fpath, None)
    
    with open(fpath, "r", encoding="utf-8") as f:
        assert f.read() == "Originale"


@patch('time.sleep', return_value=None)
def test_retry_with_backoff(mock_sleep, mock_openai):
    mock_client, mock_module = mock_openai
    
    mock_client.chat.completions.create.side_effect = [
        mock_module.RateLimitError("Too fast"),
        MagicMock(choices=[MagicMock(message=MagicMock(content="Success"))])
    ]
    
    res = writing_assistant._call_deepseek_with_retry(mock_client, [])
    assert res == "Success"
    
    mock_client.chat.completions.create.side_effect = mock_module.RateLimitError("insufficient_quota")
    with pytest.raises(RuntimeError, match="Quota API DeepSeek esaurita"):
        writing_assistant._call_deepseek_with_retry(mock_client, [])


def test_run_invalid_path():
    res = writing_assistant.run({"file_path": "/fake/path.md"})
    assert "non riesco a trovare" in res


def test_run_success_text(temp_dir, mock_openai):
    """Test end-to-end con file testo."""
    fpath = os.path.join(temp_dir, "test.md")
    with open(fpath, "w", encoding="utf-8") as f:
        f.write("Testo da migliorare.")
        
    res = writing_assistant.run({"file_path": fpath})
    assert "Ho migliorato il file" in res
    with open(fpath, "r", encoding="utf-8") as f:
        assert f.read() == "Testo migliorato di test."


def test_run_success_docx(temp_dir, mock_openai):
    """Test end-to-end con file docx che verifica in-place saving e backup."""
    from docx import Document
    fpath = os.path.join(temp_dir, "test.docx")
    doc = Document()
    doc.add_paragraph("Testo .docx da migliorare.")
    doc.save(fpath)
    
    res = writing_assistant.run({"file_path": fpath})
    assert "preservandone stili" in res
    
    # Verifica che il file docx sia stato aggiornato
    mod_doc = Document(fpath)
    assert mod_doc.paragraphs[0].text == "Testo migliorato di test."
    
    # Verifica che il backup esista
    backups = [f for f in os.listdir(temp_dir) if f.startswith("test.docx.bak_")]
    assert len(backups) == 1
